"""Read whole workbooks (every tab) from Excel / CSV / Google Sheets as raw grids."""
from __future__ import annotations

import csv
import os
import re
from pathlib import Path

Grid = list[list]


def _is_gsheet(source: str) -> bool:
    return source.startswith("http") or ("/" not in source and "." not in source)


def parse_google_url(source: str) -> tuple[str, str | None]:
    """'https://docs.google.com/spreadsheets/d/<id>/edit?gid=123#gid=123' -> (id, '123'). Bare IDs pass through."""
    m = re.search(r"/d/([A-Za-z0-9_-]+)", source)
    g = re.search(r"[?#&]gid=(\d+)", source)
    return (m.group(1) if m else source, g.group(1) if g else None)


def _gspread():
    import gspread

    return gspread.service_account(filename=os.environ["GOOGLE_SERVICE_ACCOUNT_FILE"])


def _drive_download_xlsx(file_id: str) -> bytes:
    """Fallback for Excel files stored in Drive (the Sheets API refuses them) and for any Google Sheet."""
    from google.auth.transport.requests import AuthorizedSession
    from google.oauth2.service_account import Credentials

    creds = Credentials.from_service_account_file(os.environ["GOOGLE_SERVICE_ACCOUNT_FILE"],
                                                  scopes=["https://www.googleapis.com/auth/drive.readonly"])
    sess = AuthorizedSession(creds)
    base = f"https://www.googleapis.com/drive/v3/files/{file_id}"
    r = sess.get(base, params={"alt": "media"})
    if r.status_code in (400, 403) and "export" in r.text.lower():      # native Google Sheet
        r = sess.get(f"{base}/export", params={
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
    r.raise_for_status()
    return r.content


def _xlsx_bytes_to_grids(data: bytes) -> dict[str, Grid]:
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def read_workbook(source: str) -> dict[str, Grid]:
    """tab name -> grid of cell values (row 0 = sheet row 1). Values are the displayed/calculated values."""
    if _is_gsheet(source):
        fid, _ = parse_google_url(source)
        try:
            gc = _gspread()
            sh = gc.open_by_key(fid)
            return {ws.title: ws.get_all_values() for ws in sh.worksheets()}
        except Exception as e:
            # An uploaded .xlsx opened via a Sheets URL: use the Drive API instead of the Sheets API.
            if "not supported for this document" in str(e) or "failedPrecondition" in str(e):
                return _xlsx_bytes_to_grids(_drive_download_xlsx(fid))
            raise
    ext = Path(source).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(source, data_only=True, read_only=True)
        return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    if ext == ".csv":
        with open(source, newline="", encoding="utf-8-sig") as f:
            return {Path(source).stem: list(csv.reader(f))}
    raise ValueError(f"Unsupported source: {source}")


def tab_for_gid(source: str, gid: str) -> str | None:
    """Tab title for a Google `gid`. Only reliable for native Google Sheets; None for Excel files in Drive."""
    try:
        fid, _ = parse_google_url(source)
        return _gspread().open_by_key(fid).get_worksheet_by_id(int(gid)).title
    except Exception:
        return None


def render_grid(grid: Grid, start_row: int = 1, end_row: int = 40, max_cols: int = 20, cell_width: int = 45) -> str:
    """Compact text view with 1-based row numbers and A1-style column letters, for the LLM."""
    def col(i):
        s = ""
        i += 1
        while i:
            i, r = divmod(i - 1, 26)
            s = chr(65 + r) + s
        return s

    out = []
    for n in range(start_row, min(end_row, len(grid)) + 1):
        cells = grid[n - 1][:max_cols]
        if not any(str(c).strip() for c in cells if c is not None):
            continue
        out.append(f"{n}: " + " | ".join(f"{col(i)}={str(c)[:cell_width]}" for i, c in enumerate(cells)
                                         if c is not None and str(c).strip() != ""))
    return "\n".join(out) or "(empty)"


def grid_to_rows(grid: Grid, header_row: int = 1) -> list[dict]:
    """Rows below `header_row` (1-based) as dicts keyed by header text, with `_row` = sheet row number."""
    header = [str(h).strip() for h in grid[header_row - 1]]
    rows = []
    for n, r in enumerate(grid[header_row:], start=header_row + 1):
        if not any(str(c).strip() for c in r if c is not None):
            continue
        d = {h: ("" if v is None else v) for h, v in zip(header, r) if h}
        d["_row"] = n
        rows.append(d)
    return rows


def _safe_cell(v):
    """Numbers stay numbers; text that looks like a formula is stored as text."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") and not v.replace(".", "", 1).lstrip("+-").isdigit():
        return "'" + v
    return v


def write_tabs(source: str, tabs: dict[str, Grid]) -> str:
    """Add (or replace) our own tabs next to the workpaper; original tabs are never modified.
    Google Sheet -> same spreadsheet (service account needs Editor). Excel -> a *_netsuite_ready.xlsx copy.
    CSV -> sibling CSV files. Returns a human-readable location."""
    if _is_gsheet(source):
        sh = _gspread().open_by_key(parse_google_url(source)[0])
        existing = {w.title: w for w in sh.worksheets()}
        for name, grid in tabs.items():
            ws = existing.get(name) or sh.add_worksheet(name, rows=max(len(grid) + 20, 100), cols=max(len(grid[0]), 26))
            ws.clear()
            if ws.row_count < len(grid):
                ws.resize(rows=len(grid) + 20)
            ws.update(values=[[_safe_cell(c) for c in r] for r in grid], range_name="A1", value_input_option="RAW")
            ws.freeze(rows=1)
        return f"Google Sheet tabs written: {list(tabs)}"
    p = Path(source)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(source, keep_vba=p.suffix.lower() == ".xlsm")
        for name, grid in tabs.items():
            if name in wb.sheetnames:
                del wb[name]
            ws = wb.create_sheet(name)
            for r in grid:
                ws.append([_safe_cell(c) for c in r])
            ws.freeze_panes = "A2"
        out = p.with_name(p.stem + "_netsuite_ready" + p.suffix)
        wb.save(out)
        return f"Saved copy with new tabs {list(tabs)}: {out}"
    outs = []
    for name, grid in tabs.items():
        out = p.with_name(f"{p.stem}_{name.lower().replace(' ', '_')}.csv")
        with open(out, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows([[_safe_cell(c) for c in r] for r in grid])
        outs.append(str(out))
    return "Wrote: " + ", ".join(outs)
